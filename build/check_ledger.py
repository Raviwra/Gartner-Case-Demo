"""Reconcile every figure in the spec's numbers ledger against the corpus.

Usage:
    python build/check_ledger.py          # checks corpus/*.md, writes corpus/LEDGER_CHECK.md
    python build/check_ledger.py --docx   # also checks corpus/docx/*.docx and appends the result

Each ledger entry names a document, a location, and the exact figure string that
must appear there. Table entries name the row label and the column header, and the
check reads the cell. Section entries look for the figure inside the named section.
Paragraph entries (D5) look inside the paragraph that starts with the marker.
"""
import sys
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
CORPUS = ROOT / "corpus"
OUT = CORPUS / "LEDGER_CHECK.md"

DOC_FILES = {
    "D1": "D1_borrower_financial_summary",
    "D2": "D2_loan_agreement_excerpt",
    "D3": "D3_prior_credit_memo_template",
    "D4": "D4_credit_policy_excerpt",
    "D5": "D5_relationship_history",
}

# (item, figure, doc, location, row, column)
# row/column set: table cell check. row only: paragraph marker check. neither: section substring check.
LEDGER = [
    # D1 Table 1
    ("Revenue FY2023", "61,400", "D1", "§1", "Revenue", "FY2023"),
    ("Revenue FY2024", "66,800", "D1", "§1", "Revenue", "FY2024"),
    ("Revenue FY2025", "72,500", "D1", "§1", "Revenue", "FY2025"),
    ("Revenue growth FY2023", "6.0%", "D1", "§1", "Revenue growth", "FY2023"),
    ("Revenue growth FY2024", "8.8%", "D1", "§1", "Revenue growth", "FY2024"),
    ("Revenue growth FY2025", "8.5%", "D1", "§1", "Revenue growth", "FY2025"),
    ("EBITDA FY2023", "8,600", "D1", "§1", "EBITDA", "FY2023"),
    ("EBITDA FY2024", "9,500", "D1", "§1", "EBITDA", "FY2024"),
    ("EBITDA FY2025", "10,400", "D1", "§1", "EBITDA", "FY2025"),
    ("EBITDA margin FY2023", "14.0%", "D1", "§1", "EBITDA margin", "FY2023"),
    ("EBITDA margin FY2024", "14.2%", "D1", "§1", "EBITDA margin", "FY2024"),
    ("EBITDA margin FY2025", "14.3%", "D1", "§1", "EBITDA margin", "FY2025"),
    ("Net income FY2023", "3,100", "D1", "§1", "Net income", "FY2023"),
    ("Net income FY2024", "3,600", "D1", "§1", "Net income", "FY2024"),
    ("Net income FY2025 (table)", "3,900", "D1", "§1", "Net income", "FY2025"),
    ("Unfunded capex FY2023", "1,200", "D1", "§1", "Unfunded capex", "FY2023"),
    ("Unfunded capex FY2024", "1,300", "D1", "§1", "Unfunded capex", "FY2024"),
    ("Unfunded capex FY2025", "1,400", "D1", "§1", "Unfunded capex", "FY2025"),
    ("Cash taxes FY2023", "1,000", "D1", "§1", "Cash taxes", "FY2023"),
    ("Cash taxes FY2024", "1,150", "D1", "§1", "Cash taxes", "FY2024"),
    ("Cash taxes FY2025", "1,250", "D1", "§1", "Cash taxes", "FY2025"),
    ("Debt service FY2023", "4,500", "D1", "§1", "Debt service (P+I)", "FY2023"),
    ("Debt service FY2024", "4,550", "D1", "§1", "Debt service (P+I)", "FY2024"),
    ("Debt service FY2025", "4,800", "D1", "§1", "Debt service (P+I)", "FY2025"),
    # D1 Table 2
    ("Cash FY2023", "2,100", "D1", "§2", "Cash", "FY2023"),
    ("Cash FY2024", "2,600", "D1", "§2", "Cash", "FY2024"),
    ("Cash FY2025 (deliberate gap)", "Not provided*", "D1", "§2", "Cash", "FY2025"),
    ("Revolver availability FY2023", "1,800", "D1", "§2", "Revolver availability", "FY2023"),
    ("Revolver availability FY2024", "2,300", "D1", "§2", "Revolver availability", "FY2024"),
    ("Revolver availability FY2025", "1,800", "D1", "§2", "Revolver availability", "FY2025"),
    ("Liquidity FY2023", "3,900", "D1", "§2", "Liquidity (cash + availability)", "FY2023"),
    ("Liquidity FY2024", "4,900", "D1", "§2", "Liquidity (cash + availability)", "FY2024"),
    ("Liquidity FY2025 (follows from gap)", "n/a*", "D1", "§2", "Liquidity (cash + availability)", "FY2025"),
    ("Total funded debt FY2023", "22,400", "D1", "§2", "Total funded debt", "FY2023"),
    ("Total funded debt FY2024", "20,900", "D1", "§2", "Total funded debt", "FY2024"),
    ("Total funded debt FY2025", "19,800", "D1", "§2", "Total funded debt", "FY2025"),
    ("Total leverage FY2023", "2.60x", "D1", "§2", "Total leverage", "FY2023"),
    ("Total leverage FY2024", "2.20x", "D1", "§2", "Total leverage", "FY2024"),
    ("Total leverage FY2025", "1.90x", "D1", "§2", "Total leverage", "FY2025"),
    ("DSCR FY2023", "1.42x", "D1", "§2", "DSCR", "FY2023"),
    ("DSCR FY2024", "1.55x", "D1", "§2", "DSCR", "FY2024"),
    ("DSCR FY2025", "1.61x", "D1", "§2", "DSCR", "FY2025"),
    ("Table 2 footnote: cash not in spread export", "not included in the spread export", "D1", "§2", None, None),
    ("Table 2 footnote: Q3 2026 package", "Q3 2026 package", "D1", "§2", None, None),
    # D1 Table 3 pro forma
    ("Pro forma total funded debt", "26,200", "D1", "§3", "Total funded debt", "Pro forma (FY2025 basis)"),
    ("Pro forma leverage", "2.52x", "D1", "§3", "Total leverage", "Pro forma (FY2025 basis)"),
    ("Pro forma debt service", "5,600", "D1", "§3", "Debt service", "Pro forma (FY2025 basis)"),
    ("Pro forma DSCR", "1.38x", "D1", "§3", "DSCR", "Pro forma (FY2025 basis)"),
    ("Pro forma composition: term 12,000", "term loan 12,000", "D1", "§3", None, None),
    ("Pro forma composition: revolver 3,200", "revolver 3,200", "D1", "§3", None, None),
    ("Pro forma composition: equipment notes 11,000", "equipment notes 11,000", "D1", "§3", None, None),
    ("Request: $18.0M total", "$18.0M", "D1", "§3", None, None),
    ("Request: $12.0M term loan", "$12.0M term loan", "D1", "§3", None, None),
    ("Request: $6.0M revolver", "$6.0M revolver", "D1", "§3", None, None),
    ("Use of proceeds: refinance 5,600", "term loan 5,600", "D1", "§3", None, None),
    ("Use of proceeds: expansion 6,000", "expansion 6,000", "D1", "§3", None, None),
    ("Use of proceeds: fees 400", "fees 400", "D1", "§3", None, None),
    ("Use of proceeds: total 12,000", "total of 12,000", "D1", "§3", None, None),
    # D1 §4 debt composition
    ("Debt composition: Meridian term loan", "5,600", "D1", "§4", "Meridian term loan", "FY2025"),
    ("Debt composition: Meridian revolver", "3,200", "D1", "§4", "Meridian revolver", "FY2025"),
    ("Debt composition: equipment notes", "11,000", "D1", "§4", "Equipment notes (Prairie Equipment Finance)", "FY2025"),
    ("Debt composition: total", "19,800", "D1", "§4", "Total funded debt", "FY2025"),
    # D1 §5 concentration
    ("Customer A FY2024", "28% of FY2024 revenue", "D1", "§5", None, None),
    ("Customer A FY2025", "31% of FY2025 revenue", "D1", "§5", None, None),
    ("Top five customers FY2025", "64% of FY2025 revenue", "D1", "§5", None, None),
    # D1 §6 management commentary
    ("Commentary: revenue growth 8.5%", "Revenue grew 8.5% in FY2025", "D1", "§6", None, None),
    ("Commentary: net income $4.2 million (deliberate conflict)", "Net income rose to $4.2 million", "D1", "§6", None, None),
    ("Commentary: expansion capacity 2027", "beginning in 2027", "D1", "§6", None, None),
    # D1 §7 compliance
    ("Compliance certificate date", "April 28, 2026", "D1", "§7", None, None),
    ("Compliance: leverage 1.90x", "1.90x", "D1", "§7", None, None),
    ("Compliance: max 3.50x", "3.50x", "D1", "§7", None, None),
    ("Compliance: DSCR 1.61x", "1.61x", "D1", "§7", None, None),
    ("Compliance: min 1.20x", "1.20x", "D1", "§7", None, None),
    # D2
    ("D2 draft date", "Sept 22, 2026", "D2", "__header__", None, None),
    ("D2 Founders defined", "Erik Halvorsen and Maria Halvorsen", "D2", "§1.1", None, None),
    ("D2 Schedule 1.1 gap", "[Schedule 1.1 not included in this excerpt]", "D2", "§1.1", None, None),
    ("D2 term loan", "$12,000,000", "D2", "§2.1", None, None),
    ("D2 term maturity", "five years after the Closing Date", "D2", "§2.1", None, None),
    ("D2 term amortization", "seven-year schedule", "D2", "§2.1", None, None),
    ("D2 revolver", "$6,000,000", "D2", "§2.1", None, None),
    ("D2 guaranty cap", "$3,000,000", "D2", "§2.2", None, None),
    ("D2 Rockford mortgage", "Rockford, Illinois", "D2", "§2.2", None, None),
    ("D2 §7.1 max leverage", "3.25 to 1.00", "D2", "§7.1", None, None),
    ("D2 §7.1 step-down", "3.00 to 1.00", "D2", "§7.1", None, None),
    ("D2 §7.1 step-down date", "December 31, 2027", "D2", "§7.1", None, None),
    ("D2 §7.2 min DSCR", "1.25 to 1.00", "D2", "§7.2", None, None),
    ("D2 §7.3 capex cap", "$2,500,000", "D2", "§7.3", None, None),
    ("D2 §7.3 expansion carve-out", "$6,000,000", "D2", "§7.3", None, None),
    ("D2 §7.4 annual 120 days", "120 days", "D2", "§7.4", None, None),
    ("D2 §7.4 quarterly 45 days", "45 days", "D2", "§7.4", None, None),
    ("D2 §7.5 purchase-money cap", "$1,000,000", "D2", "§7.5", None, None),
    ("D2 §7.6 'meaningful portion'", "a meaningful portion", "D2", "§7.6", None, None),
    ("D2 §7.6 'actively involved'", "actively involved", "D2", "§7.6", None, None),
    # D3
    ("D3 Section 9 marker", "[U]", "D3", "Section list", "9", "Author"),
    ("D3 Section 9 left blank", "Left blank", "D3", "Section list", "9", "Author"),
    ("D3 Section 3 hybrid marker", "[H]", "D3", "Section list", "3", "Author"),
    ("D3 prior memo: $12.0 million", "$12.0 million", "D3", "§1", None, None),
    ("D3 prior memo: $7.0 million term", "$7.0 million term loan", "D3", "§1", None, None),
    ("D3 prior memo: $5.0 million revolver", "$5.0 million revolving", "D3", "§1", None, None),
    ("D3 customer since 2019", "since 2019", "D3", "§2", None, None),
    ("D3 FY2022 revenue", "57,900", "D3", "§3", None, None),
    ("D3 FY2022 EBITDA", "7,900", "D3", "§3", None, None),
    ("D3 FY2022 margin", "13.6%", "D3", "§3", None, None),
    ("D3 FY2022 total funded debt", "22,500", "D3", "§3", None, None),
    ("D3 FY2022 leverage", "2.85x", "D3", "§3", None, None),
    ("D3 FY2022 DSCR", "1.31x", "D3", "§3", None, None),
    ("D3 3.5 omitted", "[Underwriter commentary omitted from knowledge base]", "D3", "§3", None, None),
    ("D3 2023 covenant max leverage", "3.50x", "D3", "§4", None, None),
    ("D3 2023 covenant min DSCR", "1.20x", "D3", "§4", None, None),
    ("D3 exclusion note", "including the 2023 risk rating) are excluded from this knowledge base", "D3", "§4", None, None),
    # D4
    ("D4 version and date", "v2026.1, effective March 1, 2026", "D4", "__header__", None, None),
    ("D4 4.2.1 leverage cap", "3.50x", "D4", "§4.2.1", None, None),
    ("D4 4.2.2 DSCR floor", "1.25x", "D4", "§4.2.2", None, None),
    ("D4 4.2.3 commitment threshold", "$15.0 million", "D4", "§4.2.3", None, None),
    ("D4 4.2.3 leverage threshold", "3.00x", "D4", "§4.2.3", None, None),
    ("D4 4.2.4 concentration threshold", "25%", "D4", "§4.2.4", None, None),
    ("D4 4.2.5 scale", "1-10 scale", "D4", "§4.2.5", None, None),
    ("D4 4.2.5 tool prohibition", "may not assign, recommend or pre-populate a risk rating", "D4", "§4.2.5", None, None),
    ("D4 4.2.5 prototype bracket", "[Prototype assumption:", "D4", "§4.2.5", None, None),
    # D5
    ("D5 founded 1994", "Founded 1994", "D5", "¶1", None, None),
    ("D5 employees", "About 310 employees", "D5", "¶1", None, None),
    ("D5 Erik 55%", "Erik Halvorsen 55%", "D5", "¶1", None, None),
    ("D5 Maria 25%", "Maria Halvorsen 25%", "D5", "¶1", None, None),
    ("D5 Anders 20%", "Anders Halvorsen (COO) 20%", "D5", "¶1", None, None),
    ("D5 customer since 2019", "since 2019", "D5", "¶2", None, None),
    ("D5 deposits $4.1 million", "$4.1 million", "D5", "¶2", None, None),
    ("D5 2023 facility $12.0 million", "$12.0 million facility", "D5", "¶2", None, None),
    ("D5 2023 facility split", "$7.0M term loan, $5.0M revolver", "D5", "¶2", None, None),
    ("D5 approved June 2023", "June 2023", "D5", "¶2", None, None),
    ("D5 18 days late", "18 days late", "D5", "¶3", None, None),
    ("D5 waiver number", "W-25-0117", "D5", "¶3", None, None),
    ("D5 request $18.0 million", "$18.0 million", "D5", "¶4", None, None),
    ("D5 expansion 60,000 sq ft", "60,000 sq ft", "D5", "¶4", None, None),
    ("D5 expansion $6.0 million", "$6.0 million", "D5", "¶4", None, None),
    ("D5 close by November 30, 2026", "November 30, 2026", "D5", "¶4", None, None),
    ("D5 competing term sheet 10 business days", "10 business days", "D5", "¶4", None, None),
    ("D5 Erik age 64", "Erik Halvorsen (64)", "D5", "¶5", None, None),
    ("D5 succession within three years", "within three years", "D5", "¶5", None, None),
    ("D5 minority stake discussed informally", "selling a minority stake to a private equity firm", "D5", "¶5", None, None),
    ("D5 no timeline and no LOI", "no timeline and no LOI", "D5", "¶5", None, None),
    # Footer on every document
    ("D1 footer", "FICTIONAL - prepared for case exercise.", "D1", "__any__", None, None),
    ("D2 footer", "FICTIONAL - prepared for case exercise.", "D2", "__any__", None, None),
    ("D3 footer", "FICTIONAL - prepared for case exercise.", "D3", "__any__", None, None),
    ("D4 footer", "FICTIONAL - prepared for case exercise.", "D4", "__any__", None, None),
    ("D5 footer", "FICTIONAL - prepared for case exercise.", "D5", "__any__", None, None),
]


class Block:
    """A heading plus the paragraphs and tables under it, until the next heading."""

    def __init__(self, heading):
        self.heading = heading
        self.paras = []
        self.tables = []  # each table: list of rows, each row a list of cell strings

    @property
    def text(self):
        return "\n".join(self.paras)


def parse_markdown(path):
    blocks = [Block("__header__")]
    lines = path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("#"):
            blocks.append(Block(line.lstrip("#").strip()))
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"-+", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks[-1].tables.append(rows)
            continue
        elif line.strip():
            blocks[-1].paras.append(line.replace("**", ""))
        i += 1
    return blocks


def parse_docx(path):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    blocks = [Block("__header__")]
    for child in doc.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            style = p.style.name if p.style is not None else ""
            if style.startswith("Heading") or style == "Title":
                blocks.append(Block(p.text.strip()))
            elif p.text.strip():
                blocks[-1].paras.append(p.text)
        elif tag == "tbl":
            t = Table(child, doc)
            rows = [[c.text.strip() for c in r.cells] for r in t.rows]
            blocks[-1].tables.append(rows)
    footer_texts = []
    for section in doc.sections:
        footer_texts.append("\n".join(p.text for p in section.footer.paragraphs))
    blocks.append(Block("__footer__"))
    blocks[-1].paras.extend(footer_texts)
    return blocks


def find_block(blocks, loc):
    if loc == "__header__":
        # The title line is block 1 in Markdown (title heading) and the body intro is attached to it
        return [b for b in blocks if b.heading == "__header__" or blocks.index(b) == 1]
    if loc == "__any__":
        return blocks
    out = []
    for b in blocks:
        h = b.heading
        if h == loc or h.startswith(loc + " ") or h.startswith(loc + ",") or h.startswith(loc + ":"):
            out.append(b)
    return out


def check_entry(blocks, entry):
    item, figure, doc, loc, row, col = entry
    found = blocks if loc.startswith("¶") else find_block(blocks, loc)
    if not found:
        return False, f"location {loc} not found"
    if row is not None and col is not None:
        for b in found:
            for table in b.tables:
                header = table[0]
                if col not in header:
                    continue
                ci = header.index(col)
                for r in table[1:]:
                    if r and r[0] == row:
                        cell = r[ci] if ci < len(r) else ""
                        if cell == figure or figure in cell:
                            return True, f"cell = {cell!r}"
                        return False, f"cell = {cell!r}"
        return False, f"row {row!r} / column {col!r} not found"
    if row is not None:
        # paragraph marker (D5 style): figure must sit inside the paragraph that starts with the marker
        for b in found:
            for p in b.paras:
                if p.lstrip().startswith(row) and figure in p:
                    return True, "in paragraph"
        return False, f"paragraph {row!r} with figure not found"
    if loc.startswith("¶"):
        for b in found:
            for p in b.paras:
                if p.lstrip().startswith(loc) and figure in p:
                    return True, "in paragraph"
        return False, f"paragraph {loc} with figure not found"
    for b in found:
        if figure in b.text or figure in b.heading:
            return True, "in section text"
        for table in b.tables:
            for r in table:
                if any(figure in c for c in r):
                    return True, "in section table"
    return False, "figure not in section"


def run(fmt):
    results = []
    cache = {}
    for entry in LEDGER:
        doc = entry[2]
        if doc not in cache:
            if fmt == "md":
                cache[doc] = parse_markdown(CORPUS / f"{DOC_FILES[doc]}.md")
            else:
                cache[doc] = parse_docx(CORPUS / "docx" / f"{DOC_FILES[doc]}.docx")
        ok, note = check_entry(cache[doc], entry)
        results.append((entry, ok, note))
    return results


def render(results, title, fmt_label):
    lines = [f"## {title}", ""]
    lines.append("| # | Item | Figure | Document | Location | Result | Note |")
    lines.append("|---|---|---|---|---|---|---|")
    for n, (entry, ok, note) in enumerate(results, 1):
        item, figure, doc, loc, row, col = entry
        where = loc
        if row is not None and col is not None:
            where = f"{loc}, row {row}, column {col}"
        elif row is not None:
            where = f"{loc}, {row}"
        if loc == "__header__":
            where = "header"
        if loc == "__any__":
            where = "footer (any page)"
        lines.append(
            f"| {n} | {item} | {figure.replace('|', 'pipe')} | {doc} | {where} | "
            f"{'PASS' if ok else 'FAIL'} | {note} |"
        )
    passed = sum(1 for _, ok, _ in results if ok)
    lines.append("")
    lines.append(f"**{fmt_label}: {passed} of {len(results)} checks PASS.**")
    lines.append("")
    return "\n".join(lines)


def main():
    docx_mode = "--docx" in sys.argv
    if not docx_mode:
        results = run("md")
        body = [
            "# Ledger reconciliation",
            "",
            "Every figure in the numbers ledger of the spec, where it appears in the corpus, and whether it was found exactly as written.",
            "Generated by build/check_ledger.py. Figures are $ in thousands unless the figure string says otherwise.",
            "",
            render(results, "Part 1: Markdown corpus (corpus/*.md)", "Markdown"),
        ]
        OUT.write_text("\n".join(body), encoding="utf-8")
    else:
        results = run("docx")
        existing = OUT.read_text(encoding="utf-8") if OUT.exists() else "# Ledger reconciliation\n\n"
        existing = existing.split("## Part 2:")[0].rstrip() + "\n\n"
        OUT.write_text(existing + render(results, "Part 2: Word corpus (corpus/docx/*.docx)", "Word"), encoding="utf-8")
    failed = [(e, note) for e, ok, note in results if not ok]
    print(f"{len(results) - len(failed)} of {len(results)} checks PASS ({'docx' if docx_mode else 'markdown'})")
    for e, note in failed:
        print(f"FAIL: {e[0]} -> expected {e[1]!r} in {e[2]} {e[3]} ({note})")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
