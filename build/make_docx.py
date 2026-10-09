"""Convert each corpus Markdown file (D1 to D5) into a clean .docx in corpus/docx/.

Rules:
- The first line (# heading) becomes the document title, which contains the document ID.
- "## " headings become Heading 1, "### " headings become Heading 2.
- Pipe tables become real Word tables (Table Grid style) with a bold header row.
- "- " lines become bullets. **bold** spans become bold runs.
- The line "FICTIONAL - prepared for case exercise." is removed from the body and placed in the
  page footer of every page.

Usage: python build/make_docx.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
CORPUS = ROOT / "corpus"
OUT = CORPUS / "docx"
FOOTER = "FICTIONAL - prepared for case exercise."

FILES = [
    "D1_borrower_financial_summary",
    "D2_loan_agreement_excerpt",
    "D3_prior_credit_memo_template",
    "D4_credit_policy_excerpt",
    "D5_relationship_history",
]


def add_inline(paragraph, text):
    """Write text into a paragraph, turning **bold** spans into bold runs."""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci, cell_text in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = ""
            para = cell.paragraphs[0]
            run = para.add_run(cell_text)
            if ri == 0:
                run.bold = True
            if ri > 0 and ci > 0 and re.fullmatch(r"[\d,.%x*/an ]+|Not provided\*|n/a\*", cell_text):
                para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph()


def convert(name):
    src = CORPUS / f"{name}.md"
    lines = src.read_text(encoding="utf-8").splitlines()
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Footer on every page (default footer of the single section)
    for section in doc.sections:
        footer_para = section.footer.paragraphs[0]
        footer_para.text = FOOTER
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    i = 0
    title_done = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped == FOOTER:
            i += 1
            continue
        if line.startswith("# ") and not title_done:
            doc.add_heading(line[2:].strip(), level=0)
            title_done = True
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"-+", c) for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        elif line.startswith("- "):
            para = doc.add_paragraph(style="List Bullet")
            add_inline(para, line[2:].strip())
        else:
            para = doc.add_paragraph()
            add_inline(para, stripped)
        i += 1

    OUT.mkdir(parents=True, exist_ok=True)
    dest = OUT / f"{name}.docx"
    doc.save(str(dest))
    return dest


def main():
    for name in FILES:
        dest = convert(name)
        print(f"wrote {dest.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
