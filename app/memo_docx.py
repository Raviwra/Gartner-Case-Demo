"""Export a memo (app/data/memos/<id>.json) to a Word document.

Layout:
- Title "Meridian Trust Bank, Commercial Credit Memo (Template v1)", the memo name, a small "Borrower case: ..."
  line (the memo's case; memos made before cases existed are Larkspur), and a header-field table.
- A page "Underwriter to complete (consolidated)" built from the bullets under "Underwriter to complete" in the
  machine sections. If the memo has stored AI suggestions (checklist_consolidated) the page shows those, under
  the line "AI suggestions, not applied. The underwriter decides."; otherwise it shows the checklist grouped by
  code (store.group_checklist): Documents to request, Conflicts to resolve, Policy items for judgment, Sections
  to write, Other checks, each item with the sections it came from.
- Every section in template order under a Heading 1 "Section N: Title".
  Machine sections: the Markdown converted to Word (headings, paragraphs, bullets, pipe tables, bold).
  Citations stay as plain bracket text. Underwriter sections with no text get a yellow placeholder;
  Section 9 with no text gets its own yellow placeholder. Text the underwriter typed carries a light-blue
  "Underwriter-authored" label.
- Comments become real Word comments when python-docx has Document.add_comment (1.2 or later),
  anchored to the section heading; otherwise a final "Review comments" section lists them.
- The FICTIONAL footer on every page.

This module never writes a rating, score or recommendation. Section 9 holds only what the person typed.
"""
import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

import store

TITLE = "Meridian Trust Bank, Commercial Credit Memo (Template v1)"
FOOTER = "FICTIONAL - prepared for case exercise. All documents, people and figures are invented."
HEADER_LABELS = [
    ("borrower", "Borrower"),
    ("request_type", "Request type"),
    ("total_commitment", "Total commitment"),
    ("underwriter", "Underwriter"),
    ("senior_reviewer", "Senior reviewer"),
    ("date", "Date"),
]
LIGHT_BLUE = "DEEAF6"
GARTNER_BLUE = RGBColor(0x00, 0x28, 0x69)
MUTED = RGBColor(0x4F, 0x56, 0x72)

# w:pPr children that must come after w:shd (schema order), so the shading element lands in a valid spot.
_PPR_AFTER_SHD = ("w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                  "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
                  "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                  "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                  "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def comments_supported(doc=None):
    return hasattr(doc if doc is not None else Document(), "add_comment")


# ---------------------------------------------------------------- small helpers

def _shade(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p_pr.insert_element_before(shd, *_PPR_AFTER_SHD)


def _placeholder(doc, text):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = True
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return para


def _label(doc, text, fill=LIGHT_BLUE):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = GARTNER_BLUE
    _shade(para, fill)
    return para


def _note(doc, text):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    return para


_INLINE_RE = re.compile(r"(\*\*.+?\*\*|__.+?__|`[^`]+`)")


def add_inline(paragraph, text):
    """Write Markdown inline text into a paragraph: **bold** and __bold__ become bold runs, `code`
    loses its backticks, single *emphasis* markers are dropped. Bracket citations stay as typed."""
    text = re.sub(r"(?<![*\w])\*(?!\s|\*)([^*\n]+?)(?<!\s)\*(?![*\w])", r"\1", text)
    for part in _INLINE_RE.split(text):
        if not part:
            continue
        if (part.startswith("**") and part.endswith("**") and len(part) > 4) or \
                (part.startswith("__") and part.endswith("__") and len(part) > 4):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            paragraph.add_run(part[1:-1])
        else:
            paragraph.add_run(part)


def _cells(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def add_table(doc, rows):
    width = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci in range(width):
            para = table.cell(ri, ci).paragraphs[0]
            add_inline(para, row[ci] if ci < len(row) else "")
            if ri == 0:
                for run in para.runs:
                    run.bold = True
    doc.add_paragraph()


_BULLET_RE = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(?:\[[ xX]\]\s+)?(.*)$")


def add_markdown(doc, md, base_level=2):
    """Convert section Markdown to Word. '#'/'##'/'###' headings map to Heading base_level, deeper ones
    one level lower (capped at 4). Returns the first paragraph written, or None."""
    first = None

    def mark(p):
        nonlocal first
        if first is None:
            first = p
        return p

    lines = str(md or "").replace("\r\n", "\n").split("\n")
    i = 0
    buf = []

    def flush():
        if buf:
            para = mark(doc.add_paragraph())
            add_inline(para, " ".join(s.strip() for s in buf))
            buf.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            flush()
            i += 1
            continue
        h = re.match(r"^(#{1,6})\s+(.*)$", s)
        if h:
            flush()
            n = len(h.group(1))
            level = base_level if n <= 3 else min(4, base_level + 1)
            mark(doc.add_heading(re.sub(r"[*_`]", "", h.group(2)).strip(), level=level))
            i += 1
            continue
        if s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = _cells(lines[i])
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c) or not any(cells):
                    rows.append(cells)
                i += 1
            rows = [r for r in rows if any(r)]
            if rows:
                add_table(doc, rows)
                mark(doc.paragraphs[-1])
            continue
        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", s):
            flush()
            i += 1
            continue
        b = _BULLET_RE.match(line)
        if b:
            flush()
            indent = len(b.group(1).expandtabs(4))
            numbered = b.group(2)[0].isdigit()
            style = ("List Number" if numbered else "List Bullet") + (" 2" if indent >= 2 else "")
            try:
                para = doc.add_paragraph(style=style)
            except KeyError:
                para = doc.add_paragraph(style="List Bullet")
            item = b.group(3)
            # continuation lines indented under the bullet
            while i + 1 < len(lines) and re.match(r"^\s{2,}\S", lines[i + 1]) and not _BULLET_RE.match(lines[i + 1]):
                i += 1
                item += " " + lines[i].strip()
            add_inline(mark(para), item)
            i += 1
            continue
        if s.startswith(">"):
            flush()
            para = mark(doc.add_paragraph())
            add_inline(para, s.lstrip("> ").strip())
            para.paragraph_format.left_indent = Pt(18)
            i += 1
            continue
        buf.append(s)
        i += 1
    flush()
    return first


def _fmt_time(iso):
    return str(iso or "").replace("T", " ")[:16]


# ---------------------------------------------------------------- the consolidated checklist page

def _sections_label(ids):
    ids = list(ids or [])
    return ("Section " if len(ids) == 1 else "Sections ") + ", ".join(str(i) for i in ids) if ids else ""


def _grouped_page(doc, grouped):
    """Layer 1: the checklist grouped by code (store.group_checklist). No model involved."""
    rm = grouped["removed"]
    extra = []
    if rm["subheadings"]:
        extra.append(f"{rm['subheadings']} sub-heading{'s' if rm['subheadings'] != 1 else ''} removed")
    if rm["duplicates"]:
        extra.append(f"{rm['duplicates']} repeat{'s' if rm['duplicates'] != 1 else ''} merged")
    _note(doc, f"{grouped['original_count']} items, {grouped['count']} after grouping"
               + (f" ({', '.join(extra)})" if extra else "") + ". Grouped by fixed text rules; no AI involved.")
    for g in grouped["groups"]:
        if not g["items"]:
            continue
        doc.add_heading(f"{g['title']} ({len(g['items'])})", level=2)
        for it in g["items"]:
            para = doc.add_paragraph(style="List Bullet")
            para.add_run(_sections_label(it["sections"]) + ": ").bold = True
            add_inline(para, it["text"])
            if it.get("also_cites"):
                para.add_run(" Also cited in the repeats: " + " ".join(it["also_cites"]))


def _consolidated_page(doc, cc):
    """Layer 2: the stored model suggestions (memo["checklist_consolidated"]), clearly labelled as not applied."""
    _label(doc, "AI suggestions, not applied. The underwriter decides.", fill="F2F4F7")
    _note(doc, f"Suggested by {cc.get('model') or 'the draft model'} on {_fmt_time(cc.get('at'))} from "
               f"{cc.get('source_item_count', 0)} grouped items ({cc.get('original_item_count', 0)} in the memo). "
               "Nothing here has been applied to the memo text.")

    def group(title, rows, main, extra=None, quiet=None):
        if not rows:
            return
        doc.add_heading(f"{title} ({len(rows)})", level=2)
        for r in rows:
            para = doc.add_paragraph(style="List Bullet")
            if main == "section":
                sec = str(r.get("section") or "").strip()
                para.add_run((sec if sec.lower().startswith("section") else f"Section {sec}") + ": ").bold = True
                add_inline(para, r.get(extra, ""))
                continue
            add_inline(para, r.get(main, ""))
            if extra and r.get(extra[1]):
                para.add_run(f" {extra[0]}: ").bold = True
                add_inline(para, r[extra[1]])
            if quiet and r.get(quiet[1]):
                q = para.add_run(f"\n{quiet[0]}: {r[quiet[1]]}")
                q.italic = True
                q.font.size = Pt(9)
                q.font.color.rgb = MUTED

    group("Documents to request", cc.get("documents_to_request"), "item", ("Would be in", "would_be_in"),
          ("Request line", "request_line"))
    group("Conflicts to resolve", cc.get("conflicts"), "item", ("How to resolve", "how_to_resolve"))
    group("Policy items for judgment", cc.get("policy_items"), "item", ("Judgment needed", "judgment_needed"))
    group("Sections to write", cc.get("sections_to_write"), "section", "what_it_needs")
    group("Suggested edits to the drafted sections", cc.get("suggested_edits"), "section", "suggestion")
    dropped = cc.get("dropped") or []
    if dropped:
        doc.add_heading(f"Dropped as duplicate or filler: {len(dropped)} items", level=3)
        for d in dropped:
            para = doc.add_paragraph(style="List Bullet")
            para.add_run(f"Item {d.get('item')} ({d.get('reason') or 'no reason given'}): ").bold = True
            add_inline(para, d.get("text", ""))


# ---------------------------------------------------------------- build

def build(memo):
    """Return the .docx bytes for a memo dict."""
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    for section in doc.sections:
        fp = section.footer.paragraphs[0]
        fp.text = FOOTER
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in fp.runs:
            run.font.size = Pt(9)

    doc.add_heading(TITLE, level=0)
    sub = doc.add_paragraph()
    r = sub.add_run(memo.get("name") or "")
    r.font.color.rgb = MUTED
    case = store.case_or_default(store.record_case_id(memo))
    cp = doc.add_paragraph()
    cr = cp.add_run(f"Borrower case: {case['name']} ({case['type']}). Synthetic case file.")
    cr.font.size = Pt(9)
    cr.font.color.rgb = MUTED

    header = memo.get("header") or {}
    table = doc.add_table(rows=len(HEADER_LABELS), cols=2)
    table.style = "Table Grid"
    for ri, (key, label) in enumerate(HEADER_LABELS):
        lp = table.cell(ri, 0).paragraphs[0]
        lp.add_run(label).bold = True
        vp = table.cell(ri, 1).paragraphs[0]
        val = (header.get(key) or "").strip()
        if val:
            vp.add_run(val)
        else:
            run = vp.add_run("Underwriter to complete")
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.add_paragraph()

    # Consolidated "Underwriter to complete" page, before Section 1
    doc.add_heading("Underwriter to complete (consolidated)", level=1)
    grouped = store.group_checklist(memo)
    _note(doc, "Gathered from the \"Underwriter to complete\" lists in the machine-assembled sections below. "
               "Underwriter-authored sections (3.5, 7, 8, 9 and 10) are completed in place.")
    if memo.get("checklist_consolidated"):
        _consolidated_page(doc, memo["checklist_consolidated"])
    elif grouped["count"]:
        _grouped_page(doc, grouped)
    else:
        doc.add_paragraph("No items yet. Draft the machine-assembled sections to fill this list.")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    use_comments = comments_supported(doc)
    heading_of = {}
    for sec in memo.get("sections", []):
        sid, title, author = sec["id"], sec["title"], sec.get("author")
        heading = doc.add_heading(f"Section {sid}: {title}", level=1)
        heading_of[sid] = heading
        text = (sec.get("text") or "").strip()
        if author == "M":
            if not text:
                _placeholder(doc, f"Not yet drafted: {title}. Machine-assembled section; draft it in the Memo view.")
            else:
                if sec.get("status") == "edited":
                    _label(doc, "Machine-assembled, cited. Edited by the underwriter.")
                else:
                    _label(doc, "Machine-assembled, cited", fill="F2F4F7")
                add_markdown(doc, text)
                meta = [m for m in (sec.get("model"), _fmt_time(sec.get("drafted_at"))) if m]
                if meta:
                    _note(doc, "Drafted with " + " on ".join(meta) + ". Verify citations before use.")
            if sid == "3":
                sub3 = (sec.get("text") or "")
                if not re.search(r"^#{1,6}\s*3\.5\b", sub3, re.M):
                    doc.add_heading("3.5 Underwriter Commentary", level=2)
                    _placeholder(doc, "Underwriter to complete: 3.5 Underwriter Commentary.")
        elif author == "locked":
            if text:
                _label(doc, "Underwriter-authored")
                add_markdown(doc, text)
            else:
                _placeholder(doc, "Underwriter only. Required. Not drafted or pre-populated by any tool.")
        else:
            if text:
                _label(doc, "Underwriter-authored")
                add_markdown(doc, text)
            else:
                _placeholder(doc, f"Underwriter to complete: {title}.")

    # Comments
    all_comments = [(sec["id"], c) for sec in memo.get("sections", []) for c in sec.get("comments", [])]
    if all_comments:
        if use_comments:
            for sid, c in all_comments:
                runs = heading_of[sid].runs
                body = ("[Resolved] " if c.get("resolved") else "") + c.get("text", "")
                author = c.get("author") or "Underwriter"
                initials = "".join(w[0] for w in author.split()[:3]).upper()
                doc.add_comment(runs, text=body, author=author, initials=initials)
        else:
            doc.add_heading("Review comments", level=1)
            for sid, c in all_comments:
                para = doc.add_paragraph(style="List Bullet")
                para.add_run(f"Section {sid}, {c.get('author') or 'Underwriter'}, {_fmt_time(c.get('at'))}"
                             + (" (resolved)" if c.get("resolved") else "") + ": ").bold = True
                para.add_run(c.get("text", ""))

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
